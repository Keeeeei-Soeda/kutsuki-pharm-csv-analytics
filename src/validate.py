"""Phase 0: データ契約テストと DATA_ISSUES / QUESTIONS 生成。"""

from __future__ import annotations

from collections import Counter
from datetime import timedelta
from pathlib import Path
from typing import Any, Dict, List, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from src.io import (
    PHARMACY_LAT,
    PHARMACY_LON,
    PHARMACY_NAME,
    PROCESSED_DIR,
    ROOT,
    ensure_dirs,
    load_all,
    read_csv,
)

DOCS = ROOT / "docs"
FIGURES = ROOT / "reports" / "figures"


def _setup_font() -> None:
    plt.rcParams["font.family"] = "sans-serif"
    plt.rcParams["font.sans-serif"] = [
        "Hiragino Sans",
        "AppleGothic",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def check_unique(df: pd.DataFrame, col: str, expected: int | None = None) -> Dict[str, Any]:
    n = len(df)
    nunique = df[col].nunique(dropna=False)
    dup = int(n - df[col].dropna().nunique()) if df[col].isna().any() else int(n - nunique)
    # better dup count:
    dup = int(df[col].duplicated().sum())
    out = {
        "column": col,
        "rows": n,
        "nunique": int(nunique),
        "duplicates": dup,
        "ok": dup == 0 and (expected is None or n == expected),
        "expected_rows": expected,
    }
    return out


def run_integrity_checks(data: Dict[str, pd.DataFrame]) -> List[Dict[str, Any]]:
    pm = data["patient_master"]
    vh = data["visit_history"]
    cm = data["clinic_master"]
    cmp_ = data["competitor_pharmacy"]
    wp = data["weather_pollen"]
    vt = data["visit_triangle"]
    cm_rx = data["clinic_master_from_prescriptions"]

    results: List[Dict[str, Any]] = []

    for label, check in [
        ("patient_master.患者ID", check_unique(pm, "患者ID", 8885)),
        ("visit_history.受診ID", check_unique(vh, "受診ID", 31924)),
        ("clinic_master.クリニックID", check_unique(cm, "クリニックID", 434)),
        ("competitor_pharmacy.競合薬局ID", check_unique(cmp_, "競合薬局ID", 66)),
        ("weather_pollen.日付", check_unique(wp, "日付", 706)),
    ]:
        check["name"] = label
        results.append(check)

    # weather continuity
    dates = pd.to_datetime(wp["日付"]).sort_values()
    expected_days = (dates.max() - dates.min()).days + 1
    results.append(
        {
            "name": "weather_pollen.連続日数",
            "ok": len(dates) == expected_days == 706,
            "rows": len(dates),
            "span_days": int(expected_days),
            "min": str(dates.min().date()),
            "max": str(dates.max().date()),
        }
    )

    # FK patient
    missing_patients = set(vh["患者ID"]) - set(pm["患者ID"])
    results.append(
        {
            "name": "FK visit_history.患者ID→patient_master",
            "ok": len(missing_patients) == 0,
            "missing": len(missing_patients),
        }
    )

    # FK clinic (nulls allowed = 医療機関なし)
    clinic_null = int(vh["クリニックID"].isna().sum() + (vh["クリニックID"].astype(str).str.strip() == "").sum())
    # treat empty string
    vh_cid = vh["クリニックID"].replace("", np.nan)
    clinic_null = int(vh_cid.isna().sum())
    orphan = set(vh_cid.dropna().astype(str)) - set(cm["クリニックID"].astype(str))
    results.append(
        {
            "name": "FK visit_history.クリニックID→clinic_master",
            "ok": len(orphan) == 0 and clinic_null == 6,
            "null_clinic_rows": clinic_null,
            "orphan_ids": sorted(orphan)[:10],
            "orphan_count": len(orphan),
        }
    )

    # aggregates
    sum_visits = int(pm["受診回数"].sum())
    results.append(
        {
            "name": "patient_master.受診回数合計==31924",
            "ok": sum_visits == 31924,
            "sum": sum_visits,
        }
    )

    bad_dates = int((pd.to_datetime(pm["初回受診日"]) > pd.to_datetime(pm["最終受診日"])).sum())
    results.append({"name": "初回受診日≤最終受診日", "ok": bad_dates == 0, "violations": bad_dates})

    # 受診クリニック数 vs visit distinct
    vh_cid2 = vh.copy()
    vh_cid2["クリニックID"] = vh_cid2["クリニックID"].replace("", np.nan)
    distinct = (
        vh_cid2.dropna(subset=["クリニックID"])
        .groupby("患者ID")["クリニックID"]
        .nunique()
        .reindex(pm["患者ID"])
        .fillna(0)
        .astype(int)
    )
    mismatch = int((distinct.values != pm.set_index("患者ID").loc[distinct.index, "受診クリニック数"].values).sum())
    results.append(
        {
            "name": "受診クリニック数==visit distinct",
            "ok": mismatch == 0,
            "mismatches": mismatch,
        }
    )

    # C061 missing
    ids = set(cm["クリニックID"].astype(str))
    results.append(
        {
            "name": "C061欠番（欠損扱いにしない）",
            "ok": "C061" not in ids and "C014" in ids,
            "has_C061": "C061" in ids,
            "has_C014": "C014" in ids,
        }
    )

    # clinic master count diff
    only_rx = set(cm_rx["クリニックID"]) - set(cm["クリニックID"])
    only_cm = set(cm["クリニックID"]) - set(cm_rx["クリニックID"])
    results.append(
        {
            "name": "clinic_master vs from_prescriptions 件数差",
            "clinic_master_n": len(cm),
            "from_prescriptions_n": len(cm_rx),
            "only_in_from_prescriptions": sorted(only_rx),
            "only_in_clinic_master": sorted(only_cm),
            "ok": len(cm) == 434 and len(cm_rx) == 435,
        }
    )

    # triangle vs visit ids
    results.append(
        {
            "name": "visit_triangle.受診ID == visit_history",
            "ok": set(vt["受診ID"]) == set(vh["受診ID"]),
            "triangle_n": len(vt),
            "visit_n": len(vh),
        }
    )

    return results


def compute_issue_stats(data: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    pm = data["patient_master"]
    vh = data["visit_history"]
    vt = data["visit_triangle"]
    cm = data["clinic_master"]
    vhw = data["visit_history_with_weather_pollen"]

    patient_no_addr = int(pm["緯度"].isna().sum())
    visits_no_patient_coord = int(
        vh.merge(pm[["患者ID", "緯度"]], on="患者ID", how="left")["緯度"].isna().sum()
    )
    clinic_no_coord = int(cm["緯度"].isna().sum())
    vh2 = vh.copy()
    vh2["クリニックID"] = vh2["クリニックID"].replace("", np.nan)
    m = vh2.merge(
        cm[["クリニックID", "緯度"]].rename(columns={"緯度": "クリニック緯度"}),
        on="クリニックID",
        how="left",
    )
    visits_no_clinic_coord = int(m["クリニック緯度"].isna().sum())

    coord_precision = pm["座標精度"].value_counts(dropna=False).to_dict()
    coarse = int(pm["座標精度"].isin(["町名", "市区町村"]).sum())

    pollen_missing = {c: int(vhw[c].isna().sum()) for c in ["スギ", "ヒノキ科", "花粉合計"]}
    pollen_observed_visits = int(vhw["花粉合計"].notna().sum())
    new_flag = pm["新患フラグ"].value_counts().to_dict()

    gate_str = vt["門前フラグ"].astype(str).str.upper()
    gate_true = int(gate_str.eq("TRUE").sum())
    gate_false = int(gate_str.eq("FALSE").sum())
    gate_blank = int(len(vt) - gate_true - gate_false)
    pattern = vt["立地パターン"].value_counts(dropna=False).to_dict()

    # 徒歩/自転車所要は visit_history 側（triangle には直線・道路kmのみ）
    road = pd.to_numeric(vh["患者→薬局_道路km"], errors="coerce")
    walk = pd.to_numeric(vh["患者→薬局_徒歩所要分"], errors="coerce")
    bike = pd.to_numeric(vh["患者→薬局_自転車所要分"], errors="coerce")
    mask = road.notna() & walk.notna()
    walk_resid = (walk[mask] - road[mask] / 4.8 * 60).abs().median()
    bike_resid = (bike[mask] - road[mask] / 15 * 60).abs().median()
    detour = pd.to_numeric(vt["導線迂回率"], errors="coerce")

    vh_m = vh.copy()
    vh_m["年月"] = pd.to_datetime(vh_m["受診日"]).dt.to_period("M").astype(str)
    monthly = vh_m.groupby("年月").size().rename("件数")
    trade = cm["商圏内フラグ"].astype(str).str.upper().value_counts().to_dict()

    return {
        "patient_coord_missing": patient_no_addr,
        "visits_patient_coord_missing": visits_no_patient_coord,
        "clinic_coord_missing": clinic_no_coord,
        "visits_clinic_coord_missing_or_null_id": visits_no_clinic_coord,
        "coord_precision": {str(k): int(v) for k, v in coord_precision.items()},
        "coarse_precision_patients": coarse,
        "pollen_missing_in_visit_weather": pollen_missing,
        "pollen_observed_visits": pollen_observed_visits,
        "new_patient_flag": {str(k): int(v) for k, v in new_flag.items()},
        "gate_true": gate_true,
        "gate_false": gate_false,
        "gate_blank": gate_blank,
        "location_pattern": {str(k): int(v) for k, v in pattern.items()},
        "walk_formula_median_abs_resid_min": float(walk_resid) if mask.any() else None,
        "bike_formula_median_abs_resid_min": float(bike_resid) if mask.any() else None,
        "monthly_visits": monthly.to_dict(),
        "monthly_mean": float(monthly.mean()),
        "n_months": int(len(monthly)),
        "trade_area_flag": {str(k): int(v) for k, v in trade.items()},
        "age_fixed_note": "年齢は基準日2026-07-31時点で固定（辞書記載）",
        "address_text_intentional_empty": True,
        "sunday_visits": int((pd.to_datetime(vh["受診日"]).dt.dayofweek == 6).sum()),
        "max_patient_pharmacy_km": float(pd.to_numeric(pm["患者→薬局_道路km"], errors="coerce").max()),
        "max_detour": float(detour.max()),
    }


def write_data_issues(integrity: List[Dict[str, Any]], stats: Dict[str, Any], data: Dict[str, pd.DataFrame]) -> Path:
    cm = data["clinic_master"]
    cm_rx = data["clinic_master_from_prescriptions"]
    only_rx = sorted(set(cm_rx["クリニックID"]) - set(cm["クリニックID"]))
    only_rx_names = cm_rx.set_index("クリニックID").loc[only_rx, "クリニック名"].to_dict() if only_rx else {}

    failed = [r for r in integrity if not r.get("ok", True)]
    lines = [
        "# DATA_ISSUES（Phase 0 自動生成）",
        "",
        f"- 生成対象店舗: **{PHARMACY_NAME}**（固定点 {PHARMACY_LAT}, {PHARMACY_LON}）",
        "- 受診期間: 2024-09-02 〜 2026-07-31",
        "",
        "## わかったこと（先行）",
        "",
        f"1. 門前フラグ TRUE: **{stats['gate_true']:,}** 件 / 立地パターン門前型: **{stats['location_pattern'].get('門前型', 0):,}** 件 → 実態は門前中心。",
        f"2. 月平均受診件数: **{stats['monthly_mean']:.1f}** 件（{stats['n_months']}か月）。依頼書の月1,800枚とは乖離。",
        "3. 患者住所テキストは両患者ファイルで全空欄だが、これは座標利用のための**意図的設計**（依頼者確認済）。GISは緯度経度を使用する。",
        "",
        "## 整合性テスト結果",
        "",
        f"- 全件OK: **{'YES' if not failed else 'NO'}**（失敗 {len(failed)} / {len(integrity)}）",
        "",
    ]
    for r in integrity:
        mark = "OK" if r.get("ok", True) else "NG"
        lines.append(f"- [{mark}] `{r.get('name')}` — `{ {k:v for k,v in r.items() if k!='name'} }`")

    lines += [
        "",
        "### clinic_master(434) vs clinic_master_from_prescriptions(435)",
        "",
        f"- from_prescriptions のみ: `{only_rx}` / 名称: `{only_rx_names}`",
        "- C061 は欠番（C014へ統合済）。欠損として扱わない。",
        "",
        "## 欠損・制約（実データ再計算）",
        "",
        "| 項目 | 実測 | 分析上の扱い |",
        "|---|---|---|",
        f"| 患者座標なし | {stats['patient_coord_missing']}人 / 受診{stats['visits_patient_coord_missing']}件 | GISから除外。感度分析 |",
        f"| クリニック座標なし or 医療機関なし | 施設{stats['clinic_coord_missing']} / 受診側欠損扱い{stats['visits_clinic_coord_missing_or_null_id']} | GISから除外。医療機関なし6件は明示フラグ |",
        f"| 座標精度 | {stats['coord_precision']}（町名+市区町村={stats['coarse_precision_patients']}人） | メッシュ分析で重み/除外 |",
        f"| 花粉欠損（visit_weather） | {stats['pollen_missing_in_visit_weather']} / 実測あり受診 {stats['pollen_observed_visits']} | **0埋め禁止**。飛散期のみ |",
        f"| 新患フラグ | {stats['new_patient_flag']} | **使用禁止**。初回受診日==受診日で再定義 |",
        f"| 年齢 | {stats['age_fixed_note']} | 時系列は生年月日から再計算（生年月日は現状全空欄→要追加取得） |",
        f"| 日曜受診 | {stats['sunday_visits']}件 | 曜日モデルは月〜土 |",
        f"| 徒歩/自転車所要の線形従属 | 中央残差(分) 徒歩={stats['walk_formula_median_abs_resid_min']}, 自転車={stats['bike_formula_median_abs_resid_min']} | 道路kmと同時投入禁止 |",
        f"| 遠方外れ値 | 患者→薬局道路km max={stats['max_patient_pharmacy_km']:.1f}, 迂回率 max={stats['max_detour']:.1f} | 主分析は商圏内（例≤10km） |",
        f"| 商圏内クリニック | {stats['trade_area_flag']} | 軸A/B設計に使用 |",
        f"| 住所テキスト空欄 | 意図的（座標利用） | 世帯推定は現状不可。visit_triangleの市区町村で地区分析 |",
        "",
        "## 受領時の正規化メモ",
        "",
        "- `patient_master` / `patient_clinic_with_distance` で年齢階級の Excel 日付化（`5月9日`→`5-9`, `10月14日`→`10-14`）と日付 `YYYY/M/D`→`YYYY-MM-DD` を修正済み。",
        "- `data/raw` の当該2ファイルは修正後コピーを含む。以降 raw は読み取り専用として扱う。",
        "- 生年月日は全空欄のまま（マスクの可能性）。",
        "",
        "## 禁止事項リマインダ",
        "",
        "- 花粉0埋め / 新患フラグ使用 / 徒歩・自転車・道路km同時投入 / 固定年齢を受診時点扱い / 推計を実測扱い / 自店のみで市場選択率主張 / 識別なし因果表現",
        "",
    ]
    path = DOCS / "DATA_ISSUES.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_questions(stats: Dict[str, Any]) -> Path:
    lines = [
        "# 依頼者への確認事項（QUESTIONS_TO_CLIENT）",
        "",
        "## Phase 0 で確定させたいこと",
        "",
        "1. **対象店舗**は「くつき薬局沢良宜店」か「**南茨木店**」か。データバンク固定点は南茨木店。",
        f"2. 目標の「月1,800枚」の定義は？ 実測の月平均受診明細は約 **{stats['monthly_mean']:.0f}件**（{stats['n_months']}か月）。処方箋枚数 / 受付回数 / 明細件数のどれか。",
        "3. 開局日・移転日・現住所での営業開始日は？",
        "4. 大淀店・沢良宜店も同一スキーマで入手可能か（外部妥当性・DiD）。",
        "",
        "## 優先度A（理論完成に必須）",
        "",
        "1. クリニック別の総処方箋発行数（自店シェア分母）",
        "2. 競合薬局の規模代理指標（薬剤師数、店舗面積、口コミ件数等）",
        "3. 他店舗（大淀・沢良宜）の同一スキーマデータ",
        "4. イベント年表（開局・移転・営業時間変更・近隣クリニック開閉院・競合開閉局・広告）",
        "",
        "## 優先度B",
        "",
        "5. 受診ごとの診療科・薬効分類・処方日数・点数",
        "6. 患者の性別",
        "7. 待ち時間実測と時間帯別来局数",
        "8. Googleビジネスプロフィール指標の時系列",
        "9. 祝日カレンダー（またはフラグ付与の承認）",
        "10. **生年月日**の提供可否（現状 patient_master で全空欄。受診時点年齢の再計算に必要）",
        "",
        "## 優先度C（Experience Engine）",
        "",
        "11. 患者アンケート（安心・信頼・満足・NPS・認知）300〜400件、患者ID紐付け可能な同意設計",
        "12. 初回来局のきっかけ（門前/以前から/紹介/検索/看板）",
        "",
        "## データ設計の確認（済・メモ）",
        "",
        "- 住所テキスト空欄＋座標保持は Google Maps 利用のための意図的設計である旨、依頼者確認済。",
        "",
    ]
    path = DOCS / "QUESTIONS_TO_CLIENT.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def build_marts(data: Dict[str, pd.DataFrame], stats: Dict[str, Any]) -> Dict[str, Path]:
    ensure_dirs()
    pm = data["patient_master"].copy()
    vh = data["visit_history"].copy()
    vt = data["visit_triangle"].copy()
    cm = data["clinic_master"].copy()

    vh["クリニックID"] = vh["クリニックID"].replace("", np.nan)
    vh["医療機関なしフラグ"] = vh["クリニックID"].isna()
    vh["受診年月"] = pd.to_datetime(vh["受診日"]).dt.to_period("M").astype(str)

    # join triangle pattern onto visits
    visits = vh.merge(
        vt[
            [
                "受診ID",
                "門前フラグ",
                "立地パターン",
                "年齢",
                "年齢階級",
                "患者住所_市区町村",
                "クリニック名",
                "診療科",
                "導線迂回率",
                "新患フラグ",
            ]
        ],
        on="受診ID",
        how="left",
        suffixes=("", "_tri"),
    )
    visits["軸"] = np.where(
        visits["立地パターン"].astype(str).eq("門前型")
        | visits["門前フラグ"].astype(str).str.upper().eq("TRUE"),
        "門前",
        "非門前",
    )
    # Prefer 立地パターン: 門前型 vs others
    visits["軸"] = np.where(visits["立地パターン"].astype(str).eq("門前型"), "門前", "非門前")
    visits.loc[visits["立地パターン"].isna(), "軸"] = "不明"

    patients = pm.copy()
    # first visit pattern
    first = (
        visits.sort_values("受診日")
        .groupby("患者ID", as_index=False)
        .first()[["患者ID", "立地パターン", "門前フラグ", "クリニックID", "軸"]]
        .rename(
            columns={
                "立地パターン": "初回_立地パターン",
                "門前フラグ": "初回_門前フラグ",
                "クリニックID": "初回_クリニックID",
                "軸": "初回_軸",
            }
        )
    )
    patients = patients.merge(first, on="患者ID", how="left")

    clinics = cm.copy()

    # market areas from mesh
    mesh = data["mesh_population"].copy()
    market_areas = mesh.copy()
    market_areas["距離帯_km"] = pd.cut(
        market_areas["くつき薬局南茨木店→メッシュ中心_直線km"],
        bins=[-0.01, 1, 2, 3, 5, 10],
        labels=["0-1", "1-2", "2-3", "3-5", "5+"],
    )

    paths = {}
    for name, df in [
        ("visits", visits),
        ("patients", patients),
        ("clinics", clinics),
        ("market_areas", market_areas),
    ]:
        out = PROCESSED_DIR / f"{name}.parquet"
        df.to_parquet(out, index=False)
        # also csv for easy inspection
        df.to_csv(PROCESSED_DIR / f"{name}.csv", index=False, encoding="utf-8-sig")
        paths[name] = out
    return paths


def plot_monthly(data: Dict[str, pd.DataFrame], stats: Dict[str, Any]) -> Path:
    _setup_font()
    ensure_dirs()
    vh = data["visit_history"].copy()
    vh["年月"] = pd.to_datetime(vh["受診日"]).dt.to_period("M")
    monthly = vh.groupby("年月").size()

    vt = data["visit_triangle"].copy()
    vt["年月"] = pd.to_datetime(vt["受診日"]).dt.to_period("M")
    vt["門前型"] = vt["立地パターン"].astype(str).eq("門前型")
    stacked = vt.groupby(["年月", "門前型"]).size().unstack(fill_value=0)
    stacked.columns = ["非門前型" if c is False else "門前型" for c in stacked.columns]

    fig, ax = plt.subplots(figsize=(10, 4.5))
    stacked.plot(kind="bar", stacked=True, ax=ax, color=["#4C78A8", "#F58518"], width=0.85)
    ax.axhline(1800, color="#E45756", ls="--", lw=1.2, label="依頼書 1,800")
    ax.axhline(stats["monthly_mean"], color="#54A24B", ls=":", lw=1.2, label=f"実測平均 {stats['monthly_mean']:.0f}")
    ax.set_title(f"{PHARMACY_NAME} 月別受診件数（門前型/非門前型）\n期間 2024-09〜2026-07 / N={len(vt):,}")
    ax.set_xlabel("年月")
    ax.set_ylabel("受診件数")
    ax.legend(loc="upper right")
    ax.set_xticklabels([str(x) for x in stacked.index], rotation=45, ha="right")
    fig.tight_layout()
    out = FIGURES / "phase0_monthly_visits.png"
    fig.savefig(out, dpi=150)
    plt.close(fig)
    return out


def export_data_dictionary() -> Path:
    """受領 Excel を Markdown 化（要約＋ファイル一覧）。"""
    import openpyxl

    xlsx = ROOT / "data" / "raw" / "DataBank_データ一覧kai.xlsx"
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    files = []
    current = None
    cols = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        f, col = row[0], row[1]
        if f:
            if current:
                files.append((current, cols))
            current = str(f)
            cols = 1
        elif current and col:
            cols += 1
    if current:
        files.append((current, cols))

    lines = [
        "# DATA_DICTIONARY（受領ExcelのMarkdown化）",
        "",
        f"原典: `data/raw/DataBank_データ一覧kai.xlsx`（シート: {wb.sheetnames[0]}）",
        "",
        "列定義の詳細（型・欠損・出典・更新方法）は原典 Excel を正とする。以下はファイル一覧。",
        "",
        "| ファイル | 列数（辞書記載） |",
        "|---|---:|",
    ]
    for name, n in files:
        lines.append(f"| `{name}` | {n} |")
    lines += [
        "",
        "## 注意",
        "",
        "- 本リポジトリの `data/raw/*.csv` は ID プレフィックスを除いた論理名。",
        "- 詳細な列定義が必要な場合は Excel または Phase 0 のマートスキーマを参照。",
        "",
    ]
    path = DOCS / "DATA_DICTIONARY.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def run_phase0() -> Dict[str, Any]:
    ensure_dirs()
    data = load_all()
    integrity = run_integrity_checks(data)
    stats = compute_issue_stats(data)
    issues_path = write_data_issues(integrity, stats, data)
    questions_path = write_questions(stats)
    dict_path = export_data_dictionary()
    mart_paths = build_marts(data, stats)
    fig_path = plot_monthly(data, stats)
    return {
        "integrity": integrity,
        "stats": stats,
        "docs": {"issues": issues_path, "questions": questions_path, "dictionary": dict_path},
        "marts": mart_paths,
        "figure": fig_path,
        "all_ok": all(r.get("ok", True) for r in integrity),
    }


if __name__ == "__main__":
    result = run_phase0()
    print("Phase0 all_ok:", result["all_ok"])
    for r in result["integrity"]:
        if not r.get("ok", True):
            print("FAIL", r)
    print("wrote", result["docs"])
    print("figure", result["figure"])
